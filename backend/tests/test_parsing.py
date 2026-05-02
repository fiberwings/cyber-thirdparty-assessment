from pathlib import Path
from tempfile import NamedTemporaryFile

import pymupdf
from docx import Document
from openpyxl import Workbook

from app.parsing import parse_document


def test_xlsx_emits_section_with_sheet_and_row_range():
    wb = Workbook()
    ws = wb.active
    ws.title = "CAIQ"
    ws.append(["Question", "Answer"])
    ws.append(["Encrypt at rest?", "Yes, AES-256"])
    ws.append(["MFA?", "Yes"])
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        wb.save(f.name)
        chunks = parse_document(Path(f.name), "caiq.xlsx")
    assert len(chunks) == 1
    c = chunks[0]
    assert "Sheet 'CAIQ'" in c.section_path
    assert "Encrypt at rest?" in c.text
    assert "AES-256" in c.text


def test_docx_section_path_tracks_heading_hierarchy():
    d = Document()
    d.add_heading("Policy", level=1)
    d.add_heading("Encryption", level=2)
    d.add_paragraph("AES-256 at rest")
    d.add_heading("Access", level=2)
    d.add_paragraph("MFA required for admins")
    with NamedTemporaryFile(suffix=".docx", delete=False) as f:
        d.save(f.name)
        chunks = parse_document(Path(f.name), "policy.docx")
    paths = [c.section_path for c in chunks]
    assert any("Policy > Encryption" in p for p in paths)
    assert any("Policy > Access" in p for p in paths)


def test_pdf_emits_chunks_with_page_numbers():
    pdf = pymupdf.open()
    p1 = pdf.new_page()
    p1.insert_text((72, 72), "Page one body", fontsize=11)
    p2 = pdf.new_page()
    p2.insert_text((72, 72), "Page two body", fontsize=11)
    with NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        pdf.save(f.name)
        pdf.close()
        chunks = parse_document(Path(f.name), "x.pdf")
    pages = sorted({c.page for c in chunks})
    assert pages == [1, 2]
