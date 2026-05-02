"""DOCX parsing with python-docx.

Produces chunks following Word's heading hierarchy. Tables are flattened to
pipe-separated rows that ride along with the surrounding section.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from app.parsing.chunker import ParsedChunk

MAX_CHARS_PER_CHUNK = 1500


def _heading_level(paragraph) -> int | None:
    style = paragraph.style.name if paragraph.style else ""
    if style.startswith("Heading "):
        try:
            return int(style.split(" ", 1)[1])
        except (ValueError, IndexError):
            return None
    if style == "Title":
        return 1
    return None


def _table_to_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows).strip()


def _walk_body(doc):
    """Yield ('paragraph', p) and ('table', t) in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}", 1)[-1]
        if tag == "p":
            for p in doc.paragraphs:
                if p._element is child:
                    yield ("paragraph", p)
                    break
        elif tag == "tbl":
            for t in doc.tables:
                if t._element is child:
                    yield ("table", t)
                    break


def parse_docx(path: Path) -> list[ParsedChunk]:
    doc = DocxDocument(str(path))
    chunks: list[ParsedChunk] = []
    ordinal = 0

    section_stack: list[str] = []  # heading path
    buf: list[str] = []
    buf_size = 0

    def section_path() -> str:
        return " > ".join(section_stack)

    def flush():
        nonlocal buf, buf_size, ordinal
        if buf:
            chunks.append(
                ParsedChunk(
                    ord=ordinal,
                    text="\n\n".join(buf).strip(),
                    page=None,
                    section_path=section_path(),
                )
            )
            ordinal += 1
            buf, buf_size = [], 0

    for kind, node in _walk_body(doc):
        if kind == "paragraph":
            text = node.text.strip()
            if not text:
                continue
            level = _heading_level(node)
            if level is not None:
                flush()
                section_stack[:] = section_stack[: level - 1]
                section_stack.append(text)
                continue
            if buf_size + len(text) > MAX_CHARS_PER_CHUNK and buf:
                flush()
            buf.append(text)
            buf_size += len(text)
        elif kind == "table":
            t_text = _table_to_text(node)
            if not t_text:
                continue
            if buf_size + len(t_text) > MAX_CHARS_PER_CHUNK and buf:
                flush()
            buf.append(t_text)
            buf_size += len(t_text)
    flush()

    # Suppress unused-import warning while keeping the import available for callers.
    _ = qn
    return chunks
